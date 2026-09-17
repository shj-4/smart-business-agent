"""
Unit tests لميزات المرحلة الرابعة (بلا شبكة):

- مقارنة الفترات (compare_periods) في run_query و format_query_result
- حدود الفترات get_comparison_ranges
- تفضيلات التقارير الدورية (ReportPref) وCRUD الخاصة بها
- منطق الاستحقاق _report_due وجدولة الإرسال periodic_report_job
- بناء نص التقرير الدوري build_periodic_summary
- التصدير الموحد generate_export_excel
- إجماليات الأشهر monthly_totals والرسم البياني generate_monthly_chart
"""

from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import bot.reminders as reminders_mod
from app.database.crud import (
    get_comparison_ranges,
    get_report_pref,
    list_report_prefs,
    mark_report_sent,
    monthly_totals,
    run_query,
    set_report_frequency,
)
from app.database.models import Transaction
from app.timeutil import now_local, now_utc, to_local_naive, to_utc_naive

USER_A = 111
USER_B = 222


@pytest.fixture
def db_env(monkeypatch):
    """بيئة db في الذاكرة مع SessionLocal معمّق في bot.reminders."""
    from app.database.db import Base

    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    SessionMaker = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    monkeypatch.setattr(reminders_mod, "SessionLocal", SessionMaker)

    yield SessionMaker

    Base.metadata.drop_all(engine)
    engine.dispose()


def _add_transaction(db, user_id, *, tx_type, amount, dt_utc, currency="ILS"):
    tx = Transaction(
        telegram_user_id=user_id,
        type=tx_type,
        amount=Decimal(str(amount)),
        currency=currency,
        description=f"{tx_type}-{amount}",
        raw_message="raw",
        created_at=dt_utc,
    )
    db.add(tx)
    db.commit()
    return tx


def _freeze_now(monkeypatch, dt_local):
    """تثبيت الساعة المحلية عند dt (بالتوقيت المحلي) — لا اعتماد على الساعة الحقيقية."""
    import app.timeutil as tu

    monkeypatch.setattr(tu, "now_local", lambda: dt_local)


# ---------- get_comparison_ranges ----------


class TestComparisonRanges:
    def test_today_bounds(self):
        ranges = get_comparison_ranges("today")
        cur_lo, cur_hi = ranges["current"]
        prev_lo, prev_hi = ranges["previous"]
        local_now = now_local()
        today_start = to_utc_naive(local_now.replace(hour=0, minute=0, second=0, microsecond=0))
        assert cur_lo == today_start
        assert prev_hi == today_start
        assert prev_lo < cur_lo

    def test_this_month_starts_first_day(self):
        ranges = get_comparison_ranges("this_month")
        cur_lo, _ = ranges["current"]
        local_now = now_local()
        expected = to_utc_naive(local_now.replace(day=1, hour=0, minute=0, second=0, microsecond=0))
        assert cur_lo == expected
        prev_lo, prev_hi = ranges["previous"]
        assert prev_hi == expected

    def test_partial_tracks_elapsed(self, monkeypatch):
        _freeze_now(monkeypatch, datetime(2026, 9, 15, 12, 0))
        ranges = get_comparison_ranges("this_month")
        assert ranges["partial"] is True
        # Sept 15 12:00 → 14.5/30 ≈ 48%
        assert 40 < ranges["elapsed_pct"] < 60

    def test_full_end_of_month_not_partial(self, monkeypatch):
        _freeze_now(monkeypatch, datetime(2026, 12, 31, 23, 59))
        ranges = get_comparison_ranges("this_month")
        assert ranges["partial"] is False
        assert ranges["elapsed_pct"] == 100

    def test_all_time_returns_none(self):
        assert get_comparison_ranges("all_time") is None


# ---------- compare_periods metric ----------


class TestComparePeriods:
    def test_compares_current_vs_previous_month(self, db_session, monkeypatch):
        # تثبيت الساعة عند 2026-09-15 12:00 محليًا → "this_month" = [01/09, 15/09] UTC محول
        _freeze_now(monkeypatch, datetime(2026, 9, 15, 12, 0))
        cur_lo, _ = get_comparison_ranges("this_month")["current"]
        # معاملة هذا الشهر
        _add_transaction(
            db_session, USER_A, tx_type="expense", amount=300, dt_utc=cur_lo + timedelta(hours=1)
        )
        # معاملة الشهر الماضي (قبل بداية هذا الشهر)
        _add_transaction(
            db_session, USER_A, tx_type="expense", amount=100, dt_utc=cur_lo - timedelta(days=1)
        )
        _add_transaction(
            db_session, USER_A, tx_type="income", amount=50, dt_utc=cur_lo - timedelta(days=1)
        )

        result = run_query(
            db_session,
            USER_A,
            {"metric": "compare_periods", "period": "this_month", "person": None},
        )
        assert result["metric"] == "compare_periods"
        assert result["kind"] == "comparison"
        cur = result["result"]["current"]
        prev = result["result"]["previous"]
        assert float(cur["expense"]["ILS"]) == 300.0
        assert float(prev["expense"]["ILS"]) == 100.0
        assert float(prev["income"]["ILS"]) == 50.0

    def test_unsupported_period(self, db_session):
        result = run_query(
            db_session,
            USER_A,
            {"metric": "compare_periods", "period": "all_time", "person": None},
        )
        assert result["error"] == "unsupported_period"

    def test_propagates_partial_flags(self, db_session, monkeypatch):
        _freeze_now(monkeypatch, datetime(2026, 9, 15, 12, 0))
        result = run_query(
            db_session,
            USER_A,
            {"metric": "compare_periods", "period": "this_month", "person": None},
        )
        assert result["partial"] is True
        assert 40 < result["elapsed_pct"] < 60


class TestComparisonFormatting:
    def test_format_comparison_mentions_lists(self, db_session, monkeypatch):
        from bot.formatters import format_query_result

        _freeze_now(monkeypatch, datetime(2026, 9, 15, 12, 0))
        cur_lo, _ = get_comparison_ranges("this_month")["current"]
        _add_transaction(
            db_session, USER_A, tx_type="expense", amount=300, dt_utc=cur_lo + timedelta(hours=1)
        )
        result = run_query(
            db_session,
            USER_A,
            {"metric": "compare_periods", "period": "this_month", "person": None},
        )
        text = format_query_result(result)
        assert "مقارنة" in text
        assert "هذا الشهر" in text
        assert "الشهر الماضي" in text

    def test_partial_period_shows_warning(self, db_session, monkeypatch):
        from bot.formatters import format_query_result

        _freeze_now(monkeypatch, datetime(2026, 9, 15, 12, 0))
        cur_lo, _ = get_comparison_ranges("this_month")["current"]
        _add_transaction(
            db_session, USER_A, tx_type="expense", amount=300, dt_utc=cur_lo + timedelta(hours=1)
        )
        result = run_query(
            db_session,
            USER_A,
            {"metric": "compare_periods", "period": "this_month", "person": None},
        )
        text = format_query_result(result)
        assert "⚠️" in text
        assert "مضى" in text

    def test_full_period_no_warning(self, db_session, monkeypatch):
        from bot.formatters import format_query_result

        _freeze_now(monkeypatch, datetime(2026, 12, 31, 23, 59))
        result = run_query(
            db_session,
            USER_B,
            {"metric": "compare_periods", "period": "this_month", "person": None},
        )
        text = format_query_result(result)
        assert "⚠️" not in text


# ---------- ReportPref CRUD ----------


class TestReportPrefCrud:
    def test_create_and_get_pref(self, db_session):
        pref = set_report_frequency(db_session, USER_A, "daily")
        assert pref.frequency == "daily"
        assert get_report_pref(db_session, USER_A) is pref

    def test_default_off_absent(self, db_session):
        assert get_report_pref(db_session, USER_A) is None
        assert list_report_prefs(db_session) == []

    def test_update_frequency_and_deliver_time(self, db_session):
        set_report_frequency(db_session, USER_A, "weekly", deliver_time="09:30")
        pref = set_report_frequency(db_session, USER_A, "monthly")
        assert pref.frequency == "monthly"
        assert pref.deliver_time == "09:30"

    def test_off_excludes_from_list(self, db_session):
        set_report_frequency(db_session, USER_A, "daily")
        set_report_frequency(db_session, USER_B, "off")
        active = list_report_prefs(db_session)
        assert [p.telegram_user_id for p in active] == [USER_A]

    def test_mark_report_sent(self, db_session):
        pref = set_report_frequency(db_session, USER_A, "daily")
        assert pref.last_sent_at is None
        mark_report_sent(db_session, pref)
        assert pref.last_sent_at is not None


# ---------- _report_due logic ----------


class _FakePref:
    def __init__(self, frequency, last_sent_at=None, deliver_time=None):
        self.frequency = frequency
        self.last_sent_at = last_sent_at
        self.deliver_time = deliver_time


class TestReportDue:
    def test_before_deliver_time_not_due(self):
        pref = _FakePref("daily", last_sent_at=None, deliver_time="23:59")
        assert not reminders_mod._report_due(pref, now_local())

    def test_daily_due_when_after_time_and_never_sent(self):
        pref = _FakePref("daily", last_sent_at=None, deliver_time="00:00")
        assert reminders_mod._report_due(pref, now_local())

    def test_daily_no_duplicate_same_day(self):
        pref = _FakePref("daily", last_sent_at=now_utc(), deliver_time="00:00")
        assert not reminders_mod._report_due(pref, now_local())

    def test_weekly_requires_first_day_of_week(self):
        from app.timeutil import first_day_of_week

        pref = _FakePref("weekly", last_sent_at=None, deliver_time="00:00")
        fd = first_day_of_week()
        today_wd = now_local().weekday()
        if today_wd != fd:
            assert not reminders_mod._report_due(pref, now_local())
        else:
            assert reminders_mod._report_due(pref, now_local())

    def test_monthly_requires_first_day(self):
        pref = _FakePref("monthly", last_sent_at=None, deliver_time="00:00")
        if now_local().day == 1:
            assert reminders_mod._report_due(pref, now_local())
        else:
            assert not reminders_mod._report_due(pref, now_local())

    def test_off_never_due(self):
        pref = _FakePref("off", last_sent_at=None, deliver_time="00:00")
        assert not reminders_mod._report_due(pref, now_local())


# ---------- periodic_report_job ----------


class TestPeriodicReportJob:
    def test_sends_report_and_marks_sent(self, db_env):
        SessionMaker = db_env
        db = SessionMaker()
        # أُرسل آخر مرة أمس → مستحق اليوم (مع وقت تسليم 00:00)
        pref = set_report_frequency(db, USER_A, "daily", deliver_time="00:00")
        mark_report_sent(db, pref)
        pref.last_sent_at = now_utc() - timedelta(days=1)
        db.commit()
        db.close()

        context = MagicMock()
        context.bot.send_message = AsyncMock()
        import asyncio

        asyncio.run(reminders_mod.periodic_report_job(context))
        context.bot.send_message.assert_called_once()

        db = SessionMaker()
        sent_pref = get_report_pref(db, USER_A)
        assert sent_pref.last_sent_at is not None
        last_date = to_local_naive(sent_pref.last_sent_at).date()
        assert last_date == now_local().date()
        db.close()

    def test_sends_nothing_when_no_prefs(self, db_env):
        context = MagicMock()
        context.bot.send_message = AsyncMock()
        import asyncio

        asyncio.run(reminders_mod.periodic_report_job(context))
        context.bot.send_message.assert_not_called()


# ---------- build_periodic_summary ----------


class TestPeriodicSummary:
    def test_contains_expenses_incomes_and_tasks(self, db_session, monkeypatch):
        from app.timeutil import to_utc_naive
        from bot.reports import build_periodic_summary

        # تثبيت الساعة المحلية عند 10:00 → "today" = [00:00, 10:00] (بالتوقيت المحلي)
        _freeze_now(monkeypatch, datetime(2026, 9, 7, 10, 0))
        lo = to_utc_naive(datetime(2026, 9, 7, 0, 0))
        _add_transaction(
            db_session, USER_A, tx_type="expense", amount=250, dt_utc=lo + timedelta(hours=2)
        )

        text = build_periodic_summary(db_session, USER_A, "daily")
        assert "المصاريف" in text
        assert "250" in text
        assert "الإيرادات" in text
        assert "قيد الانتظار" in text

    def test_unknown_frequency_falls_back_to_daily(self, db_session):
        from bot.reports import build_periodic_summary

        text = build_periodic_summary(db_session, USER_A, "nonsense")
        assert "اليومي" in text

    def test_vat_lines_when_vat_amounts_present(self, db_session, monkeypatch):
        """حقول VAT المخزّنة (#24) تظهر أخيرًا في التقرير الدوري: محصَّلة/مدفوعة
        مصنّفة بالعملة — أول استخدام تقريري فعلي لها."""
        from app.timeutil import to_utc_naive
        from bot.reports import build_periodic_summary

        _freeze_now(monkeypatch, datetime(2026, 9, 7, 10, 0))
        lo = to_utc_naive(datetime(2026, 9, 7, 0, 0))
        db_session.add(Transaction(
            telegram_user_id=USER_A,
            type="expense",
            amount=Decimal("117"),
            currency="ILS",
            description="فاتورة ضريبية",
            raw_message="فاتورة",
            vat_rate=Decimal("17.000"),
            vat_amount=Decimal("17.00"),
            created_at=lo + timedelta(hours=1),
        ))
        db_session.add(Transaction(
            telegram_user_id=USER_A,
            type="income",
            amount=Decimal("117"),
            currency="ILS",
            description="مبيع ضريبي",
            raw_message="مبيع",
            vat_rate=Decimal("17.000"),
            vat_amount=Decimal("17.00"),
            created_at=lo + timedelta(hours=2),
        ))
        db_session.commit()

        text = build_periodic_summary(db_session, USER_A, "daily")
        assert "ضريبة القيمة المضافة محصَّلة" in text
        assert "ضريبة القيمة المضافة مدفوعة" in text
        assert "17" in text

    def test_no_vat_lines_when_none_present(self, db_session, monkeypatch):
        from bot.reports import build_periodic_summary

        _freeze_now(monkeypatch, datetime(2026, 9, 7, 10, 0))
        _add_transaction(db_session, USER_A, tx_type="expense", amount=50, dt_utc=datetime(2026, 9, 7, 2, 0))
        text = build_periodic_summary(db_session, USER_A, "daily")
        assert "ضريبة القيمة المضافة" not in text


# ---------- generate_export_excel ----------


class TestExportExcel:
    def test_three_sheets_with_period_filter(self, db_session, monkeypatch):
        from openpyxl import load_workbook

        from app.database.crud import create_task
        from bot.exporters import generate_export_excel

        _freeze_now(monkeypatch, datetime(2026, 9, 15, 12, 0))
        cur_lo, _ = get_comparison_ranges("this_month")["current"]
        _add_transaction(
            db_session, USER_A, tx_type="expense", amount=50, dt_utc=cur_lo + timedelta(hours=2)
        )
        _add_transaction(
            db_session, USER_A, tx_type="expense", amount=999, dt_utc=cur_lo - timedelta(days=20)
        )
        create_task(db_session, USER_A, {"description": "مهمة للتصدير"}, raw_message="مهمة للتصدير")

        buf = generate_export_excel(db_session, USER_A, start_utc=cur_lo)
        assert isinstance(buf, BytesIO)
        buf.seek(0)
        wb = load_workbook(buf)
        assert wb.sheetnames == ["المعاملات المالية", "المهام", "الطلبيات والملاحظات"]

        ws = wb["المعاملات المالية"]
        values = [row for row in ws.iter_rows(values_only=True)]
        joined = "\n".join("|".join(str(c) if c is not None else "" for c in row) for row in values)
        # معاملة هذا الشهر ظاهرة، ومعاملة الشهر السابق (999) غير ظاهرة
        assert "50" in joined
        assert "999" not in joined


# ---------- monthly_totals و generate_monthly_chart ----------


class TestMonthlyTotals:
    def test_returns_ordered_month_labels(self, db_session):
        months = monthly_totals(db_session, USER_A, months=3)
        assert len(months) == 3
        assert months[0]["label"] <= months[1]["label"] <= months[2]["label"]
        assert all(m["by_currency"] == {} for m in months)

    def test_zero_or_negative_months_is_safe(self, db_session):
        assert monthly_totals(db_session, USER_A, months=0)  # لم ينهار بـ IndexError
        assert monthly_totals(db_session, USER_A, months=-5)[0]["label"]

    def test_settings_reject_chart_months_zero(self):
        import pydantic

        with pytest.raises(pydantic.ValidationError):
            from app.config import Settings

            Settings(chart_months=0)

    def test_aggregates_current_month(self, db_session, monkeypatch):
        _freeze_now(monkeypatch, datetime(2026, 9, 15, 12, 0))
        cur_lo, _ = get_comparison_ranges("this_month")["current"]
        _add_transaction(
            db_session, USER_A, tx_type="expense", amount=120, dt_utc=cur_lo + timedelta(hours=1)
        )
        months = monthly_totals(db_session, USER_A, months=1)
        assert len(months) == 1
        assert float(months[0]["by_currency"]["ILS"]["expense"]) == 120.0


class TestChart:
    def test_returns_none_without_data(self, db_session):
        from app.charts import generate_monthly_chart

        assert generate_monthly_chart(db_session, USER_A) is None


class TestGlobalChart:
    """الرسم العمومي للوحة التحكم (كل المستخدمين)."""

    def test_returns_none_without_data(self, db_session):
        from app.charts import generate_global_monthly_chart

        assert generate_global_monthly_chart(db_session) is None

    def test_returns_png_bytes_when_data_exists(self, db_session):
        from app.charts import generate_global_monthly_chart
        from app.timeutil import now_local, to_utc_naive

        start = to_utc_naive(now_local().replace(day=1, hour=0, minute=0, second=0, microsecond=0))
        _add_transaction(
            db_session,
            USER_A,
            tx_type="expense",
            amount=120,
            dt_utc=start + timedelta(hours=2),
        )
        buf = generate_global_monthly_chart(db_session, months=6)
        assert buf is not None
        assert buf.getvalue().startswith(b"\x89PNG")


# ---------- رسالة تأكيد /report_on ----------


class TestReportOnMessage:
    """الرسالة يجب أن تعكس التردد الفعلي لا «يوميًا» الثابت قديمًا."""

    def _run(self, arg, db_session, monkeypatch):
        import asyncio
        from types import SimpleNamespace

        import bot.handlers as handlers

        monkeypatch.setattr(handlers, "SessionLocal", lambda: db_session)
        replies = []

        class _Msg:
            async def reply_text(self, message, **kwargs):
                replies.append(message)

        update = SimpleNamespace(
            message=_Msg(), effective_user=SimpleNamespace(id=USER_A)
        )
        ctx = SimpleNamespace(args=[arg], user_data={})
        asyncio.run(handlers.report_on_command(update, ctx))
        return replies[0]

    def test_monthly_confirmation_is_consistent(self, db_session, monkeypatch):
        text = self._run("monthly", db_session, monkeypatch)
        assert "التقرير الشهري" in text
        assert "في أول كل شهر" in text
        assert "يوميًا" not in text

    def test_weekly_confirmation_is_consistent(self, db_session, monkeypatch):
        text = self._run("weekly", db_session, monkeypatch)
        assert "التقرير الأسبوعي" in text
        assert "في أول كل أسبوع" in text
        assert "يوميًا" not in text

    def test_daily_confirmation_is_consistent(self, db_session, monkeypatch):
        text = self._run("daily", db_session, monkeypatch)
        assert "التقرير اليومي" in text
        assert "كل يوم" in text


# ---------- budget_check: لا تسريب جلسات ----------


class TestBudgetCheckSessionCleanup:
    """كل جلسة قاعدة بيانات تُغلق — الجلسة الخارجية أيضًا وليس فقط الجلسات الداخلية."""

    def test_every_session_closed_on_success(self, monkeypatch):
        import bot.reminders as reminders

        created, closed = [], []
        budget_ids = [(111,), (222,)]

        class _Q:
            def order_by(self, *a, **k):
                return self

            def filter(self, *a, **k):
                return self

            def distinct(self):
                return self

            def all(self):
                return budget_ids

            def first(self):
                return None

        class _FakeSession:
            def __init__(self):
                self.closed = False
                created.append(self)

            def query(self, *a, **k):
                return _Q()

            def close(self):
                self.closed = True
                closed.append(self)

        monkeypatch.setattr(reminders, "SessionLocal", lambda: _FakeSession())

        import asyncio

        asyncio.run(reminders.budget_check(None))

        # جلسة واحدة خارجية + جلسة داخلية لكل ميزانية (معرّف) — وكلها مغلقة
        assert len(created) == 1 + len(budget_ids)
        assert len(closed) == len(created)
        assert all(s.closed for s in created)

    def test_outer_session_closed_on_query_error(self, monkeypatch):
        import bot.reminders as reminders

        closed = []

        class _Q:
            def distinct(self):
                return self

            def order_by(self, *a, **k):
                return self

            def all(self):
                raise RuntimeError("db down")

        class _FakeSession:
            def query(self, *a, **k):
                return _Q()

            def close(self):
                closed.append(self)

        monkeypatch.setattr(reminders, "SessionLocal", lambda: _FakeSession())
        import asyncio

        asyncio.run(reminders.budget_check(None))
        assert len(closed) == 1


# ---------- تجانس فترات /report و /export ----------


class TestReportExportPeriods:
    def test_resolve_period_aliases(self):
        import bot.handlers as handlers

        assert handlers.resolve_period_arg("week") == "week"
        assert handlers.resolve_period_arg("weekly") == "week"
        assert handlers.resolve_period_arg("month") == "month"
        assert handlers.resolve_period_arg("monthly") == "month"
        assert handlers.resolve_period_arg("today") == "today"
        assert handlers.resolve_period_arg("all") == "all"
        assert handlers.resolve_period_arg("") == "all"
        assert handlers.resolve_period_arg(None) == "all"
        assert handlers.resolve_period_arg("quarterly") is None

    def _run(self, coro):
        import asyncio

        return asyncio.run(coro)

    def _assert_warns_on_unknown(self, command):
        from types import SimpleNamespace

        replies = []

        class _Msg:
            async def reply_text(self, text, **kwargs):
                replies.append(text)

        update = SimpleNamespace(
            message=_Msg(), effective_user=SimpleNamespace(id=USER_A)
        )
        ctx = SimpleNamespace(args=["quarterly"], user_data={})
        self._run(command(update, ctx))
        assert replies and "لم أفهم الفترة" in replies[0]

    def test_report_unknown_period_replies_warning(self):
        import bot.handlers as handlers

        self._assert_warns_on_unknown(handlers.report_command)

    def test_export_unknown_period_replies_warning(self):
        import bot.handlers as handlers

        self._assert_warns_on_unknown(handlers.export_command)


# ---------- نطاق /budget: عملة|currency / شخص|person ----------


class TestBudgetCommandScope:
    def _run(self, args, db_session, monkeypatch):
        import asyncio
        from types import SimpleNamespace

        import bot.handlers as handlers

        monkeypatch.setattr(handlers, "SessionLocal", lambda: db_session)
        replies = []

        class _Msg:
            async def reply_text(self, text, **kwargs):
                replies.append(text)

        update = SimpleNamespace(
            message=_Msg(), effective_user=SimpleNamespace(id=USER_A)
        )
        ctx = SimpleNamespace(args=args, user_data={})
        asyncio.run(handlers.budget_command(update, ctx))
        return replies

    def test_scope_currency_english(self, db_session, monkeypatch):
        replies = self._run(["add", "currency", "ILS", "2000"], db_session, monkeypatch)
        assert replies and "تم إنشاء ميزانية" in replies[0]

    def test_scope_person_english(self, db_session, monkeypatch):
        replies = self._run(["add", "person", "محمد", "1500"], db_session, monkeypatch)
        assert replies and "تم إنشاء ميزانية" in replies[0]

    def test_scope_arabic_still_accepted(self, db_session, monkeypatch):
        replies = self._run(["add", "شخص", "محمد", "1500"], db_session, monkeypatch)
        assert replies and "تم إنشاء ميزانية" in replies[0]

    def test_unknown_scope_rejected_with_clear_message(self, db_session, monkeypatch):
        replies = self._run(["add", "supplier", "محمد", "1500"], db_session, monkeypatch)
        assert replies and "النطاق غير معروف" in replies[0]

    def test_unknown_currency_rejected_with_clear_message(self, db_session, monkeypatch):
        replies = self._run(["add", "currency", "XYZ", "2000"], db_session, monkeypatch)
        assert replies and "العملة غير معروفة" in replies[0]

    def test_scope_person_multi_word_target(self, db_session, monkeypatch):
        from app.database.crud import list_budgets

        replies = self._run(
            ["add", "person", "أبو", "محمد", "1500"], db_session, monkeypatch
        )
        assert replies and "تم إنشاء ميزانية" in replies[0]
        assert "أبو محمد" in replies[0]
        budgets = list_budgets(db_session, USER_A)
        assert not budgets or budgets[0].person == "أبو محمد"

    def test_scope_category_multi_word_target(self, db_session, monkeypatch):
        replies = self._run(
            ["add", "category", "مشتريات", "المكتب", "800"], db_session, monkeypatch
        )
        assert replies and "تم إنشاء ميزانية" in replies[0]
        from app.database.crud import list_budgets

        budgets = list_budgets(db_session, USER_A)
        target = (
            budgets[0].category
            if budgets and budgets[0].category
            else (budgets[0].person if budgets else None)
        )
        assert target in ("مشتريات المكتب", "مشتريات المكتب")

    def test_scope_person_without_amount_rejected(self, db_session, monkeypatch):
        replies = self._run(
            ["add", "person", "أبو", "محمد"], db_session, monkeypatch
        )
        assert replies and "لم أستطع قراءة" in replies[0]


# ---------- /credit: أسماء متعددة الكلمات ----------


class TestCreditCommandMultiWord:
    def _run(self, args, db_session, monkeypatch):
        import asyncio
        from types import SimpleNamespace

        import bot.handlers as handlers

        monkeypatch.setattr(handlers, "SessionLocal", lambda: db_session)
        replies = []

        class _Msg:
            async def reply_text(self, text, **kwargs):
                replies.append(text)

        update = SimpleNamespace(
            message=_Msg(), effective_user=SimpleNamespace(id=USER_A)
        )
        ctx = SimpleNamespace(args=args, user_data={})
        asyncio.run(handlers.credit_command(update, ctx))
        return replies

    def test_add_multi_word_person(self, db_session, monkeypatch):
        from app.database.crud import list_credit_limits

        replies = self._run(["add", "أبو", "محمد", "5000"], db_session, monkeypatch)
        assert replies and "حُدّد سقف ائتماني" in replies[0]
        assert "أبو محمد" in replies[0]
        persons = [lim.person for lim in list_credit_limits(db_session, USER_A)]
        assert "أبو محمد" in persons

    def test_add_single_word_still_works(self, db_session, monkeypatch):
        from app.database.crud import list_credit_limits

        replies = self._run(["add", "محمد", "5000"], db_session, monkeypatch)
        assert replies and "حُدّد سقف ائتماني" in replies[0]
        persons = [lim.person for lim in list_credit_limits(db_session, USER_A)]
        assert "محمد" in persons

    def test_add_without_amount_rejected(self, db_session, monkeypatch):
        replies = self._run(["add", "أبو", "محمد"], db_session, monkeypatch)
        assert replies and "استخدم: /credit إضافة" in replies[0]


# ---------- تنبيه الميزانيات: يصل لكل أعضاء المساحة ----------


class TestBudgetAlertBroadcast:
    def test_alert_sent_to_all_workspace_members(self, db_session, monkeypatch):
        import bot.reminders as reminders
        from app.database.crud import (
            accept_workspace_invite,
            create_budget,
            create_workspace,
            invite_to_workspace,
        )

        owner, partner = 700, 800
        create_workspace(db_session, owner)
        invite_to_workspace(db_session, owner, partner)
        accept_workspace_invite(db_session, partner, owner)
        budget = create_budget(db_session, owner, "person", "محمد", "100")
        budget.alerted_status = 0
        db_session.commit()

        sent = []

        class _FakeBot:
            async def send_message(self, chat_id, text, parse_mode=None):
                sent.append(chat_id)

        context = SimpleNamespace(bot=_FakeBot())
        usage = {"limit": "100", "spent": "150", "percent": 150, "over": True}
        import asyncio

        asyncio.run(reminders._notify_budget(context, db_session, budget, usage))

        assert set(sent) == {owner, partner}
        assert budget.alerted_status == 2
